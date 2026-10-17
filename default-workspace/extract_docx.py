import sys, json, zipfile
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

SRC = r"D:\Downloads\小红书虚拟电商 · 求职资料项目作战室.docx"
OUT = r"D:\Documents\deepseek-harness\default-workspace\docx_dump.md"


def iter_block_items(parent):
    from docx.document import Document as _Doc
    if isinstance(parent, _Doc):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)


doc = Document(SRC)
lines = []


def walk(parent, depth=0):
    for block in iter_block_items(parent):
        if isinstance(block, Paragraph):
            txt = block.text.strip()
            style = block.style.name if block.style is not None else ""
            if txt:
                prefix = ""
                if style.startswith("Heading") or style.startswith("标题"):
                    prefix = "#" * 2 + " "
                elif "List" in style or "列表" in style:
                    prefix = "- "
                lines.append(f"{prefix}{txt}   [style={style}]")
        else:
            lines.append("")
            lines.append("<<TABLE>>")
            for row in block.rows:
                cells = []
                for c in row.cells:
                    cells.append(" ".join(p.text.strip() for p in c.paragraphs if p.text.strip()))
                lines.append("| " + " | ".join(cells) + " |")
            lines.append("<<END TABLE>>")
            lines.append("")


walk(doc)

lines.append("")
lines.append("=== HEADERS/FOOTERS ===")
for i, sec in enumerate(doc.sections):
    for name, part in (("header", sec.header), ("footer", sec.footer)):
        txts = [p.text.strip() for p in part.paragraphs if p.text.strip()]
        if txts:
            lines.append(f"sec{i} {name}: {txts}")

lines.append("")
lines.append("=== IMAGES / MEDIA ===")
with zipfile.ZipFile(SRC) as z:
    for n in z.namelist():
        if n.startswith("word/media"):
            lines.append(f"{n}  {z.getinfo(n).file_size} bytes")

with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))

print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
print("wrote", OUT)
