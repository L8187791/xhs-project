"""生成《秋招网申截止时间表》Excel 骨架：2 子表 + 10 列 + 公式 + 四档条件格式。"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

OUT = r"D:\Documents\deepseek-harness\default-workspace\秋招网申截止时间表_v1骨架.xlsx"

HDR = Font(bold=True, color="FFFFFF", size=10)
HDR_FILL = PatternFill("solid", fgColor="2F5597")
TITLE = Font(bold=True, size=13)
BOLD = Font(bold=True)
NOTE = Font(size=9, color="808080")
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(wrap_text=True, vertical="center")
THIN = Side(style="thin", color="D0D0D0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()

# ============================================================
# Sheet 1 · 网申时间表
# ============================================================
ws = wb.active
ws.title = "网申时间表"

COLUMNS = [
    ("企业名称", 18), ("所属行业", 11), ("岗位/批次", 14),
    ("网申开始", 12), ("网申截止", 12), ("截止口径", 15),
    ("核验日期", 12), ("核验来源", 42), ("当前状态", 12), ("备注", 24),
]

ws["A1"] = "2027 届秋招网申截止时间表"
ws["A1"].font = TITLE
ws["A2"] = "红=今天截止 · 橙=3天内 · 绿=进行中 · 灰=已截止　｜　每行均标注核验日期与来源官网"
ws["A2"].font = NOTE

HEADER_ROW = 4
for i, (name, width) in enumerate(COLUMNS, start=1):
    c = ws.cell(row=HEADER_ROW, column=i, value=name)
    c.font, c.fill, c.alignment, c.border = HDR, HDR_FILL, CENTER, BOX
    ws.column_dimensions[get_column_letter(i)].width = width

# 预置 25 行空行：E 列填日期，I 列自动出状态
FIRST = HEADER_ROW + 1
LAST = FIRST + 24
for r in range(FIRST, LAST + 1):
    for i in range(1, 11):
        cell = ws.cell(row=r, column=i)
        cell.border = BOX
        cell.alignment = WRAP if i in (1, 3, 8, 10) else CENTER
    ws.cell(row=r, column=4).number_format = "yyyy-mm-dd"
    ws.cell(row=r, column=5).number_format = "yyyy-mm-dd"
    ws.cell(row=r, column=7).number_format = "yyyy-mm-dd"
    # I 列：状态自动计算
    ws.cell(row=r, column=9, value=(
        f'=IF($E{r}="","",'
        f'IF($E{r}<TODAY(),"已截止",'
        f'IF($E{r}=TODAY(),"今天截止",'
        f'IF($E{r}-TODAY()<=3,$E{r}-TODAY()&" 天(临期)",$E{r}-TODAY()&" 天"))))'
    )).font = BOLD

rng = f"A{FIRST}:J{LAST}"
# 四档条件格式：以 E 列（截止日期）为依据，整行标色
ws.conditional_formatting.add(rng, FormulaRule(
    formula=[f'AND($E{FIRST}<>"",$E{FIRST}<TODAY())'],
    fill=PatternFill("solid", fgColor="D9D9D9"), stopIfTrue=False))
ws.conditional_formatting.add(rng, FormulaRule(
    formula=[f'AND($E{FIRST}<>"",$E{FIRST}=TODAY())'],
    fill=PatternFill("solid", fgColor="FFC7CE"), stopIfTrue=False))
ws.conditional_formatting.add(rng, FormulaRule(
    formula=[f'AND($E{FIRST}<>"",$E{FIRST}-TODAY()>=1,$E{FIRST}-TODAY()<=3)'],
    fill=PatternFill("solid", fgColor="FFEB9C"), stopIfTrue=False))
ws.conditional_formatting.add(rng, FormulaRule(
    formula=[f'AND($E{FIRST}<>"",$E{FIRST}-TODAY()>3)'],
    fill=PatternFill("solid", fgColor="C6EFCE"), stopIfTrue=False))

# 示例行（提示格式，非真实数据 → 浅灰斜体说明）
import datetime
_ex_verify = datetime.date(2026, 10, 9)          # 核验日期（G 列）
_ex_deadline = datetime.date(2026, 10, 20)       # 网申截止（E 列，示例：未来日期 → 应显示 "X 天"）
ws.cell(row=FIRST, column=1, value="【示例行·填完请删除】某银行").font = Font(italic=True, color="808080")
ws.cell(row=FIRST, column=2, value="银行").font = Font(italic=True, color="808080")
ws.cell(row=FIRST, column=5, value=_ex_deadline).font = Font(italic=True, color="808080")
ws.cell(row=FIRST, column=5).number_format = "yyyy-mm-dd"
ws.cell(row=FIRST, column=6, value="固定日期").font = Font(italic=True, color="808080")
ws.cell(row=FIRST, column=7, value=_ex_verify).font = Font(italic=True, color="808080")
ws.cell(row=FIRST, column=7).number_format = "yyyy-mm-dd"
ws.cell(row=FIRST, column=8, value="https://（官网页面完整网址，直接复制）").font = Font(italic=True, color="808080")
ws.cell(row=FIRST, column=10, value="E列填截止日→I列自动出状态；G列=我打开官网看到的那天").font = Font(italic=True, color="808080")

ws.freeze_panes = f"A{FIRST}"
ws.auto_filter.ref = f"A{HEADER_ROW}:J{LAST}"

# ============================================================
# Sheet 2 · 使用说明
# ============================================================
ws2 = wb.create_sheet("使用说明")
ws2.column_dimensions["A"].width = 4
ws2.column_dimensions["B"].width = 96

readme = [
    ("T", "这份表怎么用"),
    ("", ""),
    ("1", "红色 = 今天截止，橙色 = 3 天内截止。打开第一眼先看这两档。"),
    ("2", "「核验日期」列 = 我打开官网亲眼看到的那天。核验日期比截止日期更值得看，因为它说明这条信息的新鲜度。"),
    ("3", "「核验来源」列是官网页面网址。怀疑哪一行，自己点开对一遍。"),
    ("4", "「截止口径」有几个类型：固定日期 / 招满即止 / 公告发布后 X 个工作日。写“招满即止”的，建议尽早投，不要等。"),
    ("5", "「当前状态」列由公式自动计算，会随日期变化。若打开文件时提示更新外部数据，选“是”。"),
    ("6", "表里没有“大概”“预计”这类词。没有官网原文支撑的信息，不会写进这张表。"),
    ("", ""),
    ("T", "关于更新"),
    ("", "每周我会重新核验一遍官网。更新版直接发给你，已购的不用再买。"),
    ("", ""),
    ("T", "关于错误"),
    ("", "如果你发现哪一行和官网不一致，请告诉我 —— 是我的核验出了问题，我会立刻改，并感谢你。"),
    ("", ""),
    ("T", "数据来源口径"),
    ("", "只采用企业官方招聘官网、官方公众号、高校就业信息网的公告。第三方汇总页（培训机构、信息聚合站）不作为来源。"),
]
r = 2
for tag, text in readme:
    if tag == "T":
        c = ws2.cell(row=r, column=2, value=text)
        c.font = Font(bold=True, size=12)
    elif text:
        ws2.cell(row=r, column=1, value=tag).font = BOLD
        ws2.cell(row=r, column=2, value=text).alignment = WRAP
    r += 1

wb.save(OUT)
print("saved", OUT)
print("data rows:", LAST - FIRST + 1)
