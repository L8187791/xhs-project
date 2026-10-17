# -*- coding: utf-8 -*-
"""网申进度台账模板 · 生成脚本（任务卡 30 · 任务 A / P0）

输出：交付_99元全流程资料包\⑦ 关闸清单（每周五更新）\网申进度台账模板.xlsx

硬要求（卡面）：
  1. 「剩余天数」是活公式（不是写死的值）
  2. 条件格式三档：≤7 天红 / 8–15 天橙 / 已过灰
  3. 第二张表「填写示例」（3–5 行）
  4. 第三张表「使用说明」：三步
  5. 顶部写「更新节奏：每周五发布新版；承诺至 2027 年 7 月」
  6. 禁词：保证不漏 / 永不遗漏 / 100% 不错过 / AI 陪跑
"""
import argparse
import datetime
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

WS = r"D:\Documents\deepseek-harness\default-workspace"
DEFAULT_OUT = os.path.join(
    WS, "交付_99元全流程资料包", "⑦ 关闸清单（每周五更新）", "网申进度台账模板.xlsx")
DEFAULT_DATA = os.path.join(WS, "技能包_时间表查询官", "data", "timetable.json")

TOP_LINE = "更新节奏：每周五发布新版；承诺至 2027 年 7 月"
DISCLAIMER = "帮你整理提醒，以官方最新为准"

HEADERS = ["序号", "企业名称", "岗位", "网申截止", "剩余天数",
           "投递日", "笔试日", "面试日", "状态", "备注"]
HEADER_ROW = 2          # 第 1 行是顶部承诺行
FIRST_DATA_ROW = 3
LAST_DATA_ROW = 202     # 200 行填写区

STATUS_OPTIONS = ["未投", "已投", "笔试", "面试", "offer", "已放弃"]

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FILL_TOP = PatternFill("solid", fgColor="FFF2E8")
FILL_HEADER = PatternFill("solid", fgColor="E8EEF7")
FILL_GRAY = PatternFill("solid", fgColor="EDEDED")
FILL_RED = PatternFill("solid", fgColor="FCE4E4")
FILL_ORANGE = PatternFill("solid", fgColor="FDEFD8")

FONT_GRAY = Font(color="9C9C9C")
FONT_RED = Font(color="C0392B", bold=True)
FONT_ORANGE = Font(color="B9770E")


def style_row(ws, row, ncols, fill=None, font=None, align="center"):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        if fill:
            cell.fill = fill
        if font:
            cell.font = font


def sheet_ledger(wb):
    ws = wb.active
    ws.title = "求职进度台账"

    # ── 第 1 行：顶部承诺行 ──
    ws.cell(row=1, column=1, value=TOP_LINE)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    style_row(ws, 1, len(HEADERS), fill=FILL_TOP, font=Font(bold=True, size=11))
    ws.row_dimensions[1].height = 26

    # ── 第 2 行：表头 ──
    for i, h in enumerate(HEADERS, start=1):
        ws.cell(row=HEADER_ROW, column=i, value=h)
    style_row(ws, HEADER_ROW, len(HEADERS), fill=FILL_HEADER, font=Font(bold=True))
    ws.row_dimensions[HEADER_ROW].height = 22

    # ── 数据区：序号 / 活公式 / 数据验证 / 日期格式 ──
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_OPTIONS) + '"',
                        allow_blank=True, showDropDown=False)
    dv.error = "请从下拉里选一个：%s" % "、".join(STATUS_OPTIONS)
    dv.errorTitle = "状态不在选项内"
    ws.add_data_validation(dv)

    for r in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        ws.cell(row=r, column=1, value=(r - FIRST_DATA_ROW + 1))
        # 活公式：网申截止(D) 是数字(日期)才算，否则留空
        ws.cell(row=r, column=5,
                value='=IF(ISNUMBER(D%d),D%d-TODAY(),"")' % (r, r))
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
        # 日期列格式
        for c in (4, 6, 7, 8):
            ws.cell(row=r, column=c).number_format = "yyyy-mm-dd"
        dv.add(ws.cell(row=r, column=9))

    # ── 条件格式三档（E 列：剩余天数）──
    rng = "E%d:E%d" % (FIRST_DATA_ROW, LAST_DATA_ROW)
    # ① 已过（负数）灰 —— 先判，命中即停
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="lessThan", formula=["0"], fill=FILL_GRAY, font=FONT_GRAY, stopIfTrue=True))
    # ② ≤7 天 红
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="between", formula=["0", "7"], fill=FILL_RED, font=FONT_RED, stopIfTrue=True))
    # ③ 8–15 天 橙
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="between", formula=["8", "15"], fill=FILL_ORANGE, font=FONT_ORANGE, stopIfTrue=True))

    # ── 列宽 ──
    widths = [6, 22, 20, 13, 10, 13, 13, 13, 10, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A%d" % FIRST_DATA_ROW

    # ── 底部边界声明 ──
    note_row = LAST_DATA_ROW + 2
    ws.cell(row=note_row, column=1, value=DISCLAIMER)
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(HEADERS))
    ws.cell(row=note_row, column=1).font = Font(color="7F7F7F", italic=True)
    return ws


def sheet_example(wb, rows):
    ws = wb.create_sheet("填写示例")
    ws.cell(row=1, column=1, value=TOP_LINE)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    style_row(ws, 1, len(HEADERS), fill=FILL_TOP, font=Font(bold=True, size=11))

    ws.cell(row=2, column=1,
            value="下面 5 行是示范填法，不是真实数据；请照着格式填你自己的。（截止日以你拿到的时间表为准）")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADERS))
    ws.cell(row=2, column=1).font = Font(color="7F7F7F", italic=True)

    for i, h in enumerate(HEADERS, start=1):
        ws.cell(row=3, column=i, value=h)
    style_row(ws, 3, len(HEADERS), fill=FILL_HEADER, font=Font(bold=True))

    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_OPTIONS) + '"', allow_blank=True)
    ws.add_data_validation(dv)

    for j, row in enumerate(rows):
        r = 4 + j
        for i, v in enumerate(row, start=1):
            ws.cell(row=r, column=i, value=v)
            ws.cell(row=r, column=i).border = BORDER
            ws.cell(row=r, column=i).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=5, value='=IF(ISNUMBER(D%d),D%d-TODAY(),"")' % (r, r))
        for c in (4, 6, 7, 8):
            ws.cell(row=r, column=c).number_format = "yyyy-mm-dd"
        dv.add(ws.cell(row=r, column=9))

    rng = "E4:E%d" % (3 + len(rows))
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="lessThan", formula=["0"], fill=FILL_GRAY, font=FONT_GRAY, stopIfTrue=True))
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="between", formula=["0", "7"], fill=FILL_RED, font=FONT_RED, stopIfTrue=True))
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="between", formula=["8", "15"], fill=FILL_ORANGE, font=FONT_ORANGE, stopIfTrue=True))

    for i, w in enumerate([6, 22, 20, 13, 10, 13, 13, 13, 10, 30], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def sheet_guide(wb):
    ws = wb.create_sheet("使用说明")
    ws.cell(row=1, column=1, value=TOP_LINE)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    style_row(ws, 1, 3, fill=FILL_TOP, font=Font(bold=True, size=11))

    lines = [
        ("三步用起来", True),
        ("① 打开时间表：先看「① 2027届秋招网申截止时间表」，挑出你要投的几家。", False),
        ("② 抄进台账：把这家的企业名称、岗位、网申截止抄到「求职进度台账」表。", False),
        ("   「剩余天数」会自动算，不用你填 —— 空着就说明这家是「招满即止」，没有写死截止日。", False),
        ("③ 每周五回来更新：我们会发布新版时间表，你回来对着改一下截止日，顺带看一眼剩余天数。", False),
        ("", False),
        ("三档颜色怎么看", True),
        ("红：还剩 7 天以内", False),
        ("橙：还剩 8–15 天", False),
        ("灰：已经过了（也会显示负数）", False),
        ("", False),
        ("「状态」怎么填", True),
        ("点单元格右边的小箭头选：未投 / 已投 / 笔试 / 面试 / offer / 已放弃", False),
        ("", False),
        ("说明", True),
        (DISCLAIMER + "。", False),
        ("本台账只帮你整理和提醒，不替你决定投哪家。", False),
    ]
    r = 3
    for text, bold in lines:
        ws.cell(row=r, column=1, value=text)
        ws.cell(row=r, column=1).font = Font(bold=True) if bold else Font()
        ws.cell(row=r, column=1).alignment = Alignment(vertical="center", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        r += 1
    ws.column_dimensions["A"].width = 100
    return ws


def pick_examples(data):
    """从 timetable.json 里挑 5 条覆盖三档 + 空值的示例（含日期真实存在）"""
    cs = {c["company"]: c for c in data["companies"]}
    want = ["招商银行", "兴业银行", "宁波银行", "腾讯", "中国工商银行"]
    rows = []
    n = 1
    for name in want:
        c = cs.get(name)
        if not c:
            continue
        dl = c.get("deadline", "")
        if dl == "招满即止" or not dl:
            cell = dl if dl else ""
            note = "招满即止，不写死截止日"
        else:
            y, m, d = [int(x) for x in dl.split("-")]
            cell = datetime.date(y, m, d)
            note = "示例"
        rows.append([n, name, "（照你投的岗位填）", cell, None, None, None, None, "未投", note])
        n += 1
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", default=DEFAULT_DATA)
    ap.add_argument("--out", default=DEFAULT_OUT)
    args = ap.parse_args()

    with open(args.data, encoding="utf-8") as f:
        data = json.load(f)

    wb = Workbook()
    sheet_ledger(wb)
    sheet_example(wb, pick_examples(data))
    sheet_guide(wb)

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    wb.save(args.out)
    print("输出：%s" % args.out)
    print("工作表：%s" % wb.sheetnames)
    print("台账列数：%d（%s）" % (len(HEADERS), " / ".join(HEADERS)))
    print("活公式示例（E3）：%s" % "=IF(ISNUMBER(D3),D3-TODAY(),\"\")")
    print("状态下拉：%d 项（%s）" % (len(STATUS_OPTIONS), "/".join(STATUS_OPTIONS)))
    print("条件格式：三档（<0 灰 / 0-7 红 / 8-15 橙）")


if __name__ == "__main__":
    main()
