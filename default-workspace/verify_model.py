import openpyxl

p = r"D:\Documents\deepseek-harness\default-workspace\月入2W模型_已算值.xlsx"
wb = openpyxl.load_workbook(p, data_only=True)
print("sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n===== {name} =====")
    for r in range(1, min(ws.max_row, 30) + 1):
        vals = []
        for c in range(1, min(ws.max_column, 9) + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                v = ""
            if isinstance(v, float):
                v = round(v, 2)
            s = str(v)
            vals.append(s[:26])
        if any(x.strip() for x in vals):
            print("  " + " | ".join(vals))
