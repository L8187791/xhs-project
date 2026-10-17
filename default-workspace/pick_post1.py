"""按今天（2026-10-09）筛出用于第 1 篇的企业，并输出可直接使用的数据。"""
import openpyxl, datetime

SRC = r"C:\Users\Administrator\Desktop\2027届秋招信息汇总表-终版.xlsx"
TODAY = datetime.date(2026, 10, 9)

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["秋招总表"]

today_rows, tomorrow_rows, week_rows = [], [], []
for r in range(2, ws.max_row + 1):
    name = ws.cell(row=r, column=2).value
    if not name:
        continue
    d = ws.cell(row=r, column=9).value
    if not isinstance(d, datetime.datetime):
        continue
    dd = d.date()
    rec = {
        "企业": name,
        "行业": ws.cell(row=r, column=3).value,
        "性质": ws.cell(row=r, column=4).value,
        "批次": ws.cell(row=r, column=5).value,
        "方向": ws.cell(row=r, column=6).value,
        "地点": ws.cell(row=r, column=7).value,
        "学历": ws.cell(row=r, column=8).value,
        "截止": dd.isoformat(),
        "入口": ws.cell(row=r, column=11).value,
        "备注": ws.cell(row=r, column=12).value,
        "核验": ws.cell(row=r, column=14).value,
        "状态": ws.cell(row=r, column=15).value,
    }
    if dd == TODAY:
        today_rows.append(rec)
    elif dd == TODAY + datetime.timedelta(days=1):
        tomorrow_rows.append(rec)
    if TODAY <= dd <= TODAY + datetime.timedelta(days=7):
        week_rows.append(rec)

def show(title, rows):
    print(f"\n===== {title}（{len(rows)} 家）=====")
    for x in rows:
        print(f"· {x['企业']}｜{x['行业']}｜{x['性质']}｜{x['批次']}")
        print(f"    截止 {x['截止']}｜投递入口：{x['入口']}")
        print(f"    备注：{x['备注']}")
        print(f"    核实状态：{x['状态']}｜核验日期：{x['核验']}")

show("今天截止（2026-10-09）", today_rows)
show("明天截止（2026-10-10）", tomorrow_rows)
print(f"\n===== 本周内（10-09 ~ 10-16）共 {len(week_rows)} 家 =====")
for x in sorted(week_rows, key=lambda z: z["截止"]):
    print(f"  {x['截止']}  {x['企业']}（{x['行业']}）")
