import openpyxl

p = r"D:\Documents\deepseek-harness\default-workspace\时间表_已算值.xlsx"
wb = openpyxl.load_workbook(p, data_only=True)
ws = wb["网申时间表"]
print("sheets:", wb.sheetnames)
print("max_row:", ws.max_row)
for r in range(4, 9):
    print(f"r{r}:", [ws.cell(row=r, column=c).value for c in range(1, 11)])
