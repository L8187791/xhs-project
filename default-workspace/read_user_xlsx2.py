"""读取用户表的第 13-15 列（数据来源/更新时间/核实状态）与「使用说明」全文。"""
import openpyxl
from collections import Counter

SRC = r"C:\Users\Administrator\Desktop\2027届秋招信息汇总表-终版.xlsx"
wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["秋招总表"]

hdr = [ws.cell(row=1, column=c).value for c in range(13, 16)]
print("第13-15列表头:", hdr)

srcs, ups, stats = Counter(), Counter(), Counter()
samples = []
for r in range(2, ws.max_row + 1):
    name = ws.cell(row=r, column=2).value
    if not name:
        continue
    s = ws.cell(row=r, column=13).value
    u = ws.cell(row=r, column=14).value
    st = ws.cell(row=r, column=15).value
    srcs[str(s)] += 1
    ups[str(u)] += 1
    stats[str(st)] += 1
    if len(samples) < 5:
        samples.append((name, s, u, st))

print("\n--- 数据来源 分布 ---")
for k, v in srcs.most_common():
    print(f"  {v:>3} 家 | {k[:70]}")
print("\n--- 更新时间 分布 ---")
for k, v in ups.most_common():
    print(f"  {v:>3} 家 | {k[:40]}")
print("\n--- 核实状态 分布 ---")
for k, v in stats.most_common():
    print(f"  {v:>3} 家 | {k[:70]}")

print("\n--- 样本行 ---")
for n, s, u, st in samples:
    print(f"  {n}: 来源={str(s)[:30]} | 更新={u} | 状态={str(st)[:40]}")

print("\n===== 使用说明 全文 =====")
ws2 = wb["使用说明"]
for r in range(1, ws2.max_row + 1):
    a = ws2.cell(row=r, column=1).value
    b = ws2.cell(row=r, column=2).value
    if a or b:
        print(f"  {a} | {b}")
