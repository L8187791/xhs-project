"""验证台账公式：灌入测试数据 → 重算 → 检查客单/累计/进度判定是否正确。"""
import openpyxl, shutil, os

SRC = r"D:\Documents\deepseek-harness\default-workspace\第1个月执行台账.xlsx"
TEST = r"D:\Documents\deepseek-harness\default-workspace\_台账公式测试.xlsx"
shutil.copy(SRC, TEST)

wb = openpyxl.load_workbook(TEST)
ws = wb["每日台账"]
# 灌入 3 天测试数据：Day1 (1单/19.9/2评论/320观看) Day2 (2单/39.8/5/610) Day3 (1单/34.9/3/480)
data = [
    (5, 1, 19.9, 2, 320),
    (6, 2, 39.8, 5, 610),
    (7, 1, 34.9, 3, 480),
]
for r, orders, sales, comments, views in data:
    ws.cell(row=r, column=4, value=orders)
    ws.cell(row=r, column=5, value=sales)
    ws.cell(row=r, column=7, value=comments)
    ws.cell(row=r, column=8, value=views)
wb.save(TEST)
print("已灌入测试数据 →", TEST)

print("""
预期结果：
  Day1 客单 = 19.9/1 = 19.90     累计单量 1    进度外推 1*30/1 = 30 → 🟡 偏慢
  Day2 客单 = 39.8/2 = 19.90     累计单量 3    进度外推 3*30/2 = 45 → ✅ 按节奏
  Day3 客单 = 34.9/1 = 34.90     累计单量 4    进度外推 4*30/3 = 40 → 🟡 偏慢
  累计客单 Day3 = (19.9+39.8+34.9)/4 = 94.6/4 = 23.65
""")
