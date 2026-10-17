# -*- coding: utf-8 -*-
"""
DB28 任务 B · 49.9 三份分行业清单 xlsx
规则（严格按 DB28 卡 §2）：基准日 2026-10-19
  进行中 = date_form=='hard' 且 deadline >= 基准日
  滚动招聘 = date_form in ('rolling','none')
  已截止 = date_form=='hard' 且 deadline < 基准日（不交付）
字段（按大纲 v2 §四）：企业名称 / 截止日期 / 剩余天数 / 企业性质 / 投递入口 / 来源可信度 / 备注
用户裁决：央国企篇只收国企央企（45 家）+ 2 家 to_verify 标「待核实」列末尾
"""
import os, json
from datetime import date, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from _clean_note import clean as clean_note  # 备注清洗（只用于交付件，不回写源数据）

ROOT = r"D:\Documents\deepseek-harness\default-workspace"
SRC = os.path.join(ROOT, "技能包_时间表查询官", "data", "timetable.json")
OUTDIR = os.path.join(ROOT, "49.9产品包-上架用", "03 交付文件")
os.makedirs(OUTDIR, exist_ok=True)

BASE = date(2026, 10, 19)

# ---------- 配色（与 19.9 美化版同色系）----------
NAVY = "1F3864"
NAVY_L = "4A6FA5"
ZEBRA = "F2F5FA"
RED_BG = "FCE4E4"
RED_TX = "C00000"
BORDER_C = "D0D7E5"
GRAY_TX = "7A8391"

thin = Side(style="thin", color=BORDER_C)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

F_TITLE = Font(name="微软雅黑", size=15, bold=True, color=NAVY)
F_SUB = Font(name="微软雅黑", size=9, color=GRAY_TX)
F_HEAD = Font(name="微软雅黑", size=10.5, bold=True, color="FFFFFF")
F_BODY = Font(name="微软雅黑", size=10, color="22262D")
F_RED = Font(name="微软雅黑", size=10, bold=True, color=RED_TX)
F_GRAY = Font(name="微软雅黑", size=10, color=GRAY_TX)

FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_ZEBRA = PatternFill("solid", fgColor=ZEBRA)
FILL_RED = PatternFill("solid", fgColor=RED_BG)

AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_L = Alignment(horizontal="left", vertical="center", wrap_text=True)

COLS = ["企业名称", "截止日期", "剩余天数", "企业性质", "投递入口", "来源可信度", "备注"]
WIDTHS = [30, 12, 9, 10, 32, 10, 46]

# ---------- 数据 ----------
def parse(s):
    try:
        y, m, d = map(int, s.split("-")); return date(y, m, d)
    except Exception:
        return None

def classify(x):
    df = x.get("date_form"); dl = (x.get("deadline") or "").strip()
    if df == "hard":
        p = parse(dl)
        return "ongoing" if (p and p >= BASE) else "closed"
    if df in ("rolling", "none"):
        return "rolling"
    return "to_verify"

def days_left(x):
    """剩余天数：从基准日 2026-10-19 起算（买家上架日）"""
    c = classify(x)
    if c == "rolling":
        return "招满即止"
    if c == "to_verify":
        return "待核实"
    p = parse(x.get("deadline") or "")
    if not p:
        return "待核实"
    n = (p - BASE).days
    return n if n >= 0 else f"{n}"

def deadline_text(x):
    c = classify(x)
    if c == "rolling":
        return "招满即止"
    if c == "to_verify":
        return "待核实"
    return x.get("deadline", "")

def trust(x):
    lv = x.get("source_level", "").strip()
    return lv if lv else "—"

def build_sheet(ws, title, sub, rows):
    ncol = len(COLS)
    last = get_column_letter(ncol)

    # 标题
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws["A1"].alignment = AL_L
    ws.row_dimensions[1].height = 34

    # 说明行
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = sub
    ws["A2"].font = F_SUB
    ws["A2"].alignment = AL_L
    ws.row_dimensions[2].height = 22

    # 表头
    for j, c in enumerate(COLS, 1):
        cell = ws.cell(row=3, column=j, value=c)
        cell.font = F_HEAD
        cell.fill = FILL_HEAD
        cell.alignment = AL_C
        cell.border = BORDER
    ws.row_dimensions[3].height = 26

    # 数据
    for i, r in enumerate(rows):
        rr = 4 + i
        vals = [r["name"], r["deadline"], r["days"], r["nature"],
                r["url"], r["trust"], r["note"]]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=rr, column=j, value=v)
            cell.border = BORDER
            cell.alignment = AL_L if j in (1, 5, 7) else AL_C
            cell.font = F_BODY
            if i % 2 == 1:
                cell.fill = FILL_ZEBRA
        # 7 天内截止 → 高亮（仅限进行中且数字型剩余天数）
        d = r["days"]
        if isinstance(d, int) and 0 <= d <= 7:
            for j in (2, 3):
                c = ws.cell(row=rr, column=j)
                c.fill = FILL_RED
                c.font = F_RED
        # 待核实 → 灰字
        if r["deadline"] == "待核实":
            ws.cell(row=rr, column=2).font = F_GRAY
            ws.cell(row=rr, column=3).font = F_GRAY

    # 列宽
    for j, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    # 冻结 + 筛选
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{last}{3 + len(rows)}"

    # 打印设置（横向 A4）
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def row_of(x):
    return {
        "name": x["company"],
        "deadline": deadline_text(x),
        "days": days_left(x),
        "nature": x.get("ownership", ""),
        "url": x.get("apply_url", ""),
        "trust": trust(x),
        # 🔴 备注清洗：只用于对外交付件，timetable.json 源数据不动
        "note": clean_note((x.get("note") or "").strip()),
    }


def main():
    d = json.load(open(SRC, encoding="utf-8"))
    comps = d["companies"]

    def pick(ind, extra_verify=False):
        main = [x for x in comps if x.get("industry") == ind
                and classify(x) in ("ongoing", "rolling")]
        main.sort(key=lambda x: (classify(x) != "ongoing", x.get("deadline") or "9999"))
        rows = [row_of(x) for x in main]
        vf = []
        if extra_verify:
            vf = [x for x in comps if x.get("industry") == ind and classify(x) == "to_verify"]
            rows += [row_of(x) for x in vf]
        return rows, len(main), len(vf)

    spec = [
        ("银行篇", "银行", "49.9产品包-银行篇.xlsx", False),
        ("互联网篇", "互联网", "49.9产品包-互联网篇.xlsx", False),
        ("央国企篇", "国企/央企", "49.9产品包-央国企篇.xlsx", True),
    ]

    report = []
    for name, ind, fname, ev in spec:
        rows, n_main, n_vf = pick(ind, ev)
        wb = Workbook()
        ws = wb.active
        ws.title = name
        n_on = sum(1 for r in rows if isinstance(r["days"], int) and r["days"] >= 0)
        title = f"2026 秋招 {name} · 网申截止清单"
        sub = (f"数据基准日 2026-10-19 ｜ 共 {len(rows)} 家 ｜ "
               f"进行中 {n_on} 家 · 滚动招聘 {sum(1 for r in rows if r['deadline'] == '招满即止')} 家"
               + (f" · 待核实 {n_vf} 家" if n_vf else "")
               + " ｜ 截止时间以企业官网为准")
        build_sheet(ws, title, sub, rows)
        p = os.path.join(OUTDIR, fname)
        wb.save(p)
        report.append((name, len(rows), p))
        print(f"saved {p}  ({len(rows)} 家)")

    print("\n=== 合计 ===")
    for n, c, p in report:
        print(f"{n}: {c} 家")
    print("总计:", sum(c for _, c, _ in report))


if __name__ == "__main__":
    main()
