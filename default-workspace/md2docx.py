# -*- coding: utf-8 -*-
"""通用 Markdown → DOCX 转换器（支持表格/标题/粗体/列表/引用/代码块/复选框）
用法: python md2docx.py <input.md> <output.docx> [--title "文档标题"]
"""
import io, os, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x66, 0x66, 0x66)
ACCENT = RGBColor(0xC0, 0x27, 0x2D)


def set_cjk(run, name="微软雅黑"):
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rPr.append(rf)
    rf.set(qn("w:ascii"), name)
    rf.set(qn("w:hAnsi"), name)
    rf.set(qn("w:eastAsia"), name)


def rich(p, text, size=10.5, bold=False, color=INK):
    """处理 **粗体** 与 `代码`"""
    text = re.sub(r"`([^`]+)`", r"\1", text)
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        b = bold
        t = part
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            b = True; t = part[2:-2]
        r = p.add_run(t)
        r.font.size = Pt(size); r.font.bold = b; r.font.color.rgb = color
        set_cjk(r, "微软雅黑")


def add_table(doc, rows, header=True):
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=cols)
    t.style = "Light Grid Accent 1"
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(cols):
            txt = row[ci] if ci < len(row) else ""
            txt = re.sub(r"`([^`]+)`", r"\1", txt)
            txt = re.sub(r"\*\*([^*]+)\*\*", r"\1", txt)
            txt = txt.replace("<br>", "\n")
            para = cells[ci].paragraphs[0]
            r = para.add_run(txt)
            r.font.size = Pt(9)
            r.font.bold = header and ri == 0
            set_cjk(r)
    return t


def build(src, dst, title=None):
    md = io.open(src, encoding="utf-8").read()
    doc = Document()
    # 页边距
    for s in doc.sections:
        s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.2)
        s.left_margin = Cm(2.4); s.right_margin = Cm(2.4)
    # 正文样式
    st = doc.styles["Normal"]
    st.font.name = "微软雅黑"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    if title:
        h = doc.add_heading(title, level=0)
        for r in h.runs:
            set_cjk(r, "微软雅黑")

    lines = md.split("\n")
    i = 0
    in_code = False
    code_buf = []
    while i < len(lines):
        ln = lines[i].rstrip()

        # 代码块
        if ln.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                r = p.add_run("\n".join(code_buf))
                r.font.size = Pt(9); r.font.name = "Consolas"
                r.font.color.rgb = GREY
                code_buf = []; in_code = False
            else:
                in_code = True
            i += 1; continue
        if in_code:
            code_buf.append(ln); i += 1; continue

        # 表格
        if ln.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i]):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows)
                doc.add_paragraph()
            continue

        # 分隔线
        if re.match(r"^\s*---+\s*$", ln):
            p = doc.add_paragraph()
            r = p.add_run("─" * 30); r.font.color.rgb = GREY; r.font.size = Pt(8)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", ln)
        if m:
            lvl = len(m.group(1))
            txt = re.sub(r"\*\*([^*]+)\*\*", r"\1", m.group(2))
            h = doc.add_heading(txt, level=min(lvl, 4))
            for r in h.runs:
                set_cjk(r, "微软雅黑")
                if lvl <= 2:
                    r.font.color.rgb = ACCENT
            i += 1; continue

        # 空行
        if not ln.strip():
            i += 1; continue

        # 引用
        if ln.strip().startswith(">"):
            txt = ln.strip().lstrip(">").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            rich(p, txt, size=10, color=RGBColor(0x44, 0x44, 0x44))
            i += 1; continue

        # 复选框
        m = re.match(r"^\s*-\s*\[( |x|X)\]\s*(.*)$", ln)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            mark = "☑ " if m.group(1).lower() == "x" else "☐ "
            r = p.add_run(mark); r.font.size = Pt(10.5); set_cjk(r)
            rich(p, m.group(2))
            i += 1; continue

        # 无序列表
        m = re.match(r"^(\s*)[-*·]\s+(.*)$", ln)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            rich(p, m.group(2))
            i += 1; continue

        # 有序列表
        m = re.match(r"^(\s*)\d+[.)]\s+(.*)$", ln)
        if m:
            p = doc.add_paragraph(style="List Number")
            rich(p, m.group(2))
            i += 1; continue

        # 普通段落
        p = doc.add_paragraph()
        rich(p, ln.strip())
        i += 1

    os.makedirs(os.path.dirname(os.path.abspath(dst)), exist_ok=True)
    doc.save(dst)
    print("SAVED %s  (%d 段)" % (dst, len(doc.paragraphs)))


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    title = None
    if "--title" in sys.argv:
        title = sys.argv[sys.argv.index("--title") + 1]
    if len(args) < 2:
        print("用法: python md2docx.py <input.md> <output.docx> [--title 标题]"); sys.exit(1)
    build(args[0], args[1], title)
