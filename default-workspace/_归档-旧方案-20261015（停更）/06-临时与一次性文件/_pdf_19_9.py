# -*- coding: utf-8 -*-
"""
DB28 · 任务 A（续）：19.9 美化版 PDF 预览导出
- 输入：2026秋招网申截止时间表_美化版.xlsx
- 输出：2026秋招网申截止时间表_预览_美化版.pdf（A4 横向 / 重复表头 / 7天内高亮）
- 用 reportlab 纯 Python 排版（本机无 Excel/LibreOffice）
"""
import os
from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Table,
                                TableStyle, Paragraph, Spacer)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase.cidfonts import UnicodeCIDFont

XLSX = "19.9产品包-上架用/03 交付文件/2026秋招网申截止时间表_美化版.xlsx"
OUT  = "19.9产品包-上架用/03 交付文件/2026秋招网申截止时间表_预览_美化版.pdf"

# --- 中文字体：优先系统 msyh.ttc，否则用 reportlab 内置 STSong-Light ---
FONT = "STSong-Light"
FONT_B = "STSong-Light"
try:
    for cand in [r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\msyh.ttf",
                 r"C:\Windows\Fonts\simhei.ttf", r"C:\Windows\Fonts\simsun.ttc"]:
        if os.path.exists(cand):
            pdfmetrics.registerFont(TTFont("MSYH", cand))
            FONT = "MSYH"; FONT_B = "MSYH"
            print("[字体]", cand)
            break
    else:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        print("[字体] 内置 STSong-Light")
except Exception as e:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    FONT = FONT_B = "STSong-Light"
    print("[字体] 回退 STSong-Light:", e)

# --- 读数据 ---
wb = load_workbook(XLSX)
ws = wb.active
HDR = 3
NCOL = ws.max_column
hdr = [ws.cell(HDR, c).value for c in range(1, NCOL + 1)]

# 只取打印用列（去掉过宽/低价值列：[等级,银行类型] 保留紧凑）
KEEP = ["序号", "企业名称", "截止日期", "状态", "剩余天数", "来源可信度",
        "企业性质", "招聘方向", "投递入口"]
kidx = [hdr.index(k) for k in KEEP if k in hdr]
cols = [hdr[i] for i in kidx]

rows = []
for r in range(HDR + 1, ws.max_row + 1):
    if ws.cell(r, 1).value is None:
        continue
    vals = []
    for i in kidx:
        v = ws.cell(r, i + 1).value
        vals.append("" if v is None else str(v))
    days = ws.cell(r, hdr.index("剩余天数") + 1).value
    rows.append((vals, days))
print("[数据] 行数 =", len(rows), "| 列 =", cols)

# --- 样式 ---
st_cell   = ParagraphStyle("c", fontName=FONT, fontSize=6.6, leading=8.2,
                           textColor=colors.HexColor("#262626"))
st_cell_c = ParagraphStyle("cc", parent=st_cell, alignment=1)
st_cell_w = ParagraphStyle("cw", parent=st_cell, textColor=colors.HexColor("#C00000"),
                           fontName=FONT_B)
st_cell_wc= ParagraphStyle("cwc", parent=st_cell_w, alignment=1)
st_url    = ParagraphStyle("u", parent=st_cell, fontSize=5.4, leading=6.4,
                           textColor=colors.HexColor("#1F4E79"), wordWrap="CJK")
st_head   = ParagraphStyle("h", fontName=FONT_B, fontSize=7.2, leading=9,
                           textColor=colors.white, alignment=1)
st_title  = ParagraphStyle("t", fontName=FONT_B, fontSize=14, leading=18,
                           textColor=colors.HexColor("#1F3864"))
st_note   = ParagraphStyle("n", fontName=FONT, fontSize=7.4, leading=10,
                           textColor=colors.HexColor("#595959"))
st_foot   = ParagraphStyle("f", fontName=FONT, fontSize=7, leading=9,
                           textColor=colors.HexColor("#808080"), alignment=1)

# --- 表格数据 ---
data = [[Paragraph(h, st_head) for h in cols]]
warn_rows = []   # 记录高亮行号（表内 index，从 1 起算表头）
for i, (vals, days) in enumerate(rows, start=1):
    line = []
    is_warn = isinstance(days, (int, float)) and 0 <= days <= 7
    is_past = isinstance(days, (int, float)) and days < 0
    for j, v in enumerate(vals):
        name = cols[j]
        if name == "投递入口" and v.startswith("http"):
            # 长 URL 折行显示（按 / 与长度断点），避免溢出
            dv = v.replace("https://", "").replace("http://", "")
            line.append(Paragraph(dv, st_url))
        elif name in ("序号", "截止日期", "状态", "剩余天数", "来源可信度", "企业性质"):
            line.append(Paragraph(v, st_cell_wc if is_warn else st_cell_c))
        else:
            line.append(Paragraph(v, st_cell_w if is_warn else st_cell))
    data.append(line)
    if is_warn:
        warn_rows.append(i)

# 列宽（可用宽度）
PW, PH = landscape(A4)
avail = PW - 2 * 12 * mm
weights = {"序号": .038, "企业名称": .235, "截止日期": .090, "状态": .070,
           "剩余天数": .066, "来源可信度": .080, "企业性质": .080,
           "招聘方向": .185, "投递入口": .156}
cw = [avail * weights.get(c, .12) for c in cols]

tbl = Table(data, colWidths=cw, repeatRows=1)
tstyle = [
    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D0D7E5")),
    ("TOPPADDING", (0, 0), (-1, -1), 2.0),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 2.0),
    ("LEFTPADDING", (0, 0), (-1, -1), 2.2),
    ("RIGHTPADDING", (0, 0), (-1, -1), 2.2),
]
for i in range(1, len(data)):
    if i % 2 == 1:
        tstyle.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F2F5FA")))
for i in warn_rows:
    tstyle.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FCE4E4")))
tbl.setStyle(TableStyle(tstyle))

# --- 文档模板（页眉/页脚）---
def on_page(canv, doc):
    canv.saveState()
    canv.setFont(FONT, 7)
    canv.setFillColor(colors.HexColor("#A6A6A6"))
    canv.drawRightString(PW - 12 * mm, 8 * mm, f"第 {doc.page} 页")
    canv.drawString(12 * mm, 8 * mm, "2026 秋招网申截止时间表 · 预览版")
    canv.restoreState()

doc = BaseDocTemplate(OUT, pagesize=landscape(A4),
                      leftMargin=12 * mm, rightMargin=12 * mm,
                      topMargin=12 * mm, bottomMargin=13 * mm,
                      title="2026秋招网申截止时间表 预览")
frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="F")
doc.addPageTemplates([PageTemplate(id="P", frames=[frame], onPage=on_page)])

flow = [
    Paragraph("2026 秋招网申截止时间表", st_title),
    Spacer(1, 4),
    Paragraph("主表 131 家 ｜ 按截止日期升序 ｜ <b>剩余天数 ≤ 7 天</b> 的行以浅红底、深红字高亮 ｜ "
              "「进行中」= 还可投递，「已截止」= 仅作历史参考 ｜ 截止时间以企业官网为准",
              st_note),
    Spacer(1, 6),
    tbl,
    Spacer(1, 5),
    Paragraph("预览版仅展示部分字段；完整版 Excel 含 15 个字段（含备注、学历、工作地点、银行类型等）。",
              st_foot),
]
doc.build(flow)
print("[保存] ->", OUT)
print("[校验] 页数:", doc.page)
