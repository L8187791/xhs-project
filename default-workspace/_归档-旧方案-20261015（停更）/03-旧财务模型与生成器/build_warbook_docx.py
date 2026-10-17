# -*- coding: utf-8 -*-
"""作战计划：Markdown 主文档 → Word。每天改 md 后重跑本脚本即可同步。"""
import io, os, re, datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

WS = r"D:\Documents\deepseek-harness\default-workspace"
SRC = os.path.join(WS, "作战计划_主文档.md")
OUT1 = os.path.join(WS, "作战计划-小红书求职资料项目-2026-10-09.docx")
OUT2 = r"D:\Downloads\作战计划-小红书求职资料项目-2026-10-09.docx"

FONT = "微软雅黑"
doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.line_spacing = 1.4; st.paragraph_format.space_after = Pt(4)
for s in doc.sections:
    s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0)
    s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)


def _f(run, size=10.5, bold=False, color=None):
    run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def rich(p, text, size=10.5, base_bold=False, color=None):
    """处理 **粗体** 与 `代码`"""
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for seg in parts:
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            _f(p.add_run(seg[2:-2]), size=size, bold=True, color=color)
        elif seg.startswith("`") and seg.endswith("`"):
            r = p.add_run(seg[1:-1]); _f(r, size=size - 0.5, bold=base_bold, color=color)
            r.font.name = "Consolas"
        else:
            _f(p.add_run(seg), size=size, bold=base_bold, color=color)


def add_table(rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol); t.style = "Table Grid"
    for ri, r in enumerate(rows):
        cells = t.add_row().cells
        for i in range(ncol):
            txt = r[i] if i < len(r) else ""
            cells[i].text = ""
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            rich(p, txt, size=9.5, base_bold=(ri == 0))
            if ri == 0:
                for run in p.runs:
                    run.font.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ---------------- 封面 ----------------
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(90)
_f(p.add_run("小红书求职资料项目"), size=26, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
_f(p.add_run("作 战 计 划"), size=30, bold=True, color=(0xA0, 0x1D, 0x1D))
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
def _cover_version():
    """封面版本号从 md 文件头解析，不再手填（2026-10-11 修：
    封面一度写死 v4.3，md 已到 v4.7 而封面不同步）。"""
    try:
        with io.open(SRC, encoding="utf-8") as fh:
            for line in fh:
                m = re.search(r"\*\*版本\s+(v[\d.]+)\*\*", line)
                if m:
                    ver = m.group(1)
                    d = re.search(r"\*\*本次更新\*\*\s+(\d{4}-\d{2}-\d{2})", line)
                    return "%s ｜ 建立 2026-10-09 ｜ 更新 %s ｜ 每日更新" % (
                        ver, d.group(1) if d else datetime.date.today().isoformat())
    except OSError:
        pass
    return "v? ｜ 建立 2026-10-09 ｜ 每日更新"


_f(p.add_run(_cover_version()), size=12)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
_f(p.add_run("账号：小L整理求职资料（秋招版）（小红书号 1163622308）"), size=11)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
# v4.8 分工（用户 2026-10-11 18:45 拍板）：不再用 v4.3 的「主运营/独立复核」旧表述
_f(p.add_run("笔记全责：WorkBuddy ｜ 审核定稿 + 独立复核：豆包 ｜ 拍板 + 平台操作：你"), size=11)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
_f(p.add_run("目标：12 月月净利保底 10,000 元，1 月冲 20,000 元"), size=11, bold=True, color=(0x2C, 0x7A, 0x57))
doc.add_page_break()

# ---------------- 正文 ----------------
lines = io.open(SRC, encoding="utf-8").read().split("\n")
i = 0
first_h1 = True
while i < len(lines):
    ln = lines[i].rstrip()

    if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-\|]+\|$", lines[i + 1].strip()):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not re.match(r"^[\s:\-]+$", "".join(cells)):
                rows.append(cells)
            i += 1
        add_table(rows)
        continue

    if ln.strip() == "---":
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        i += 1; continue

    if not ln.strip():
        i += 1; continue

    m = re.match(r"^(#{1,4})\s+(.*)$", ln)
    if m:
        level = len(m.group(1)); txt = m.group(2)
        if level == 1:
            if not first_h1:
                doc.add_page_break()
            first_h1 = False
            pp = doc.add_paragraph(); pp.paragraph_format.space_before = Pt(6); pp.paragraph_format.space_after = Pt(6)
            _f(pp.add_run(txt), size=18, bold=True, color=(0xA0, 0x1D, 0x1D))
        elif level == 2:
            pp = doc.add_paragraph(); pp.paragraph_format.space_before = Pt(12); pp.paragraph_format.space_after = Pt(4)
            _f(pp.add_run(txt), size=14, bold=True, color=(0xA0, 0x1D, 0x1D))
        else:
            pp = doc.add_paragraph(); pp.paragraph_format.space_before = Pt(8); pp.paragraph_format.space_after = Pt(3)
            _f(pp.add_run(txt), size=12, bold=True)
        i += 1; continue

    if ln.startswith("> "):
        pp = doc.add_paragraph()
        pp.paragraph_format.left_indent = Cm(0.6); pp.paragraph_format.space_before = Pt(4); pp.paragraph_format.space_after = Pt(6)
        rich(pp, ln[2:], size=11, base_bold=True, color=(0xA0, 0x1D, 0x1D))
        i += 1; continue

    if re.match(r"^[-*]\s+", ln):
        pp = doc.add_paragraph(style="List Bullet")
        pp.paragraph_format.space_after = Pt(2); pp.paragraph_format.left_indent = Cm(0.75)
        rich(pp, re.sub(r"^[-*]\s+", "", ln)); i += 1; continue

    if re.match(r"^\d+[\.、]\s+", ln):
        pp = doc.add_paragraph(style="List Number")
        pp.paragraph_format.space_after = Pt(2); pp.paragraph_format.left_indent = Cm(0.75)
        rich(pp, re.sub(r"^\d+[\.、]\s+", "", ln)); i += 1; continue

    pp = doc.add_paragraph()
    rich(pp, ln)
    i += 1

doc.save(OUT1)
print("SAVED", OUT1)
try:
    import shutil
    shutil.copyfile(OUT1, OUT2)
    print("SAVED", OUT2)
except Exception as e:
    print("COPY FAIL:", e)
print("段落数:", len(doc.paragraphs), "表格数:", len(doc.tables))
